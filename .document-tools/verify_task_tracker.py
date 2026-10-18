"""Check schedule, dependencies, workbook and tracker integrity and frontend gate."""
from pathlib import Path
from datetime import date, timedelta
from collections import Counter, defaultdict
import ast
import csv
import json
import sys

sys.path.insert(0, str(Path(__file__).parent / "deps"))
from openpyxl import load_workbook
from build_task_tracker import ready_formula

root = Path(__file__).resolve().parents[1]
folder = root / "docs/project-management"
data = json.loads((folder / "tasks.json").read_text(encoding="utf-8"))
tasks = data["tasks"]
assert len(tasks) == 40
assert data["members"] == ["mem1", "mem2"]
ids = {t["id"]: t for t in tasks}
assert len(ids) == len(tasks)
load = defaultdict(int)
for task in tasks:
    assert task["owner"] in (1, 2)
    assert 1 <= task["start_day"] <= task["end_day"] <= 20
    assert task["acceptance"] and task["deliverable"]
    load[(task["owner"], task["start_day"])] += task["hours"]
    for dependency in task["deps"]:
        assert dependency in ids, (task["id"], dependency)
        assert ids[dependency]["end_day"] <= task["start_day"], (task["id"], dependency)
assert max(load.values()) <= 6
visited, visiting = set(), set()
def visit(task_id):
    assert task_id not in visiting, f"Dependency cycle: {task_id}"
    if task_id in visited:
        return
    visiting.add(task_id)
    for dependency in ids[task_id]["deps"]:
        visit(dependency)
    visiting.remove(task_id)
    visited.add(task_id)
for task in tasks:
    visit(task["id"])
assert Counter(t["phase"] for t in tasks) == {"API": 30, "FRONTEND_LATER": 6, "QA": 4}
assert all(t["end_day"] <= 15 for t in tasks if t["phase"] == "API")
assert all(t["start_day"] >= 16 for t in tasks if t["phase"] == "FRONTEND_LATER")

wb = load_workbook(folder / "task-tracker.xlsx", data_only=False)
assert wb.sheetnames == ["Config", "Tasks", "Dashboard", "Modules"]
assert wb["Config"]["B3"].value == "mem1"
assert wb["Config"]["B4"].value == "mem2"
assert wb["Config"]["B6"].value == "NOT_APPROVED"
sheet = wb["Tasks"]
assert sheet.max_row == 41 and sheet.max_column == 21
for row, task in enumerate(tasks, 2):
    assert sheet.cell(row, 1).value == task["id"]
    assert sheet.cell(row, 15).value is False
    assert sheet.cell(row, 17).value == ready_formula(task, row, 41)
    assert '"finish","unfinish"' in sheet.cell(row, 16).value
    assert sheet.cell(row, 18).value is None
    assert sheet.cell(row, 19).value is None
    assert sheet.cell(row, 20).value is None
    assert sheet.cell(row, 5).value == f'=Config!$B${3 if task["owner"] == 1 else 4}'
assert len(sheet.data_validations.dataValidation) == 1
assert sheet.freeze_panes == "E2" and sheet.auto_filter.ref == "A1:U41"
with (folder / "tasks-import.csv").open(encoding="utf-8-sig", newline="") as stream:
    rows = list(csv.reader(stream))
assert len(rows) == 41 and all(len(row) == 21 for row in rows)
assert all(row[15] == "unfinish" for row in rows[1:])

current = date.fromisoformat(data["start_date"])
dates = []
while len(dates) < 20:
    if current.weekday() < 5:
        dates.append(current)
    current += timedelta(days=1)
assert dates[-1] == date(2026, 11, 6)
for name in ["core", "accounts", "academics", "students", "schedules", "attendance", "qr", "reports"]:
    assert (root / "backend/school_app" / name).is_dir()
# User-authorized local demo UI is an explicit exception; gate/finish stay unchanged.
ui_root = root / "backend/school_app/demo_ui"
for relative in ("templates/demo_ui/index.html", "static/demo_ui/app.css", "static/demo_ui/app.js"):
    assert (ui_root / relative).is_file(), f"Missing local demo UI: {relative}"
authorization = (root / "AGENTS.md").read_text(encoding="utf-8")
assert "làm luôn UI để demo" in authorization and "NOT_APPROVED" in authorization
assert "FRONTEND_LATER remains historical baseline" in data["scope_note"]
assert data["version"] == "1.3"
for row, task in zip(rows[1:], tasks):
    assert row[0] == task["id"] and row[3] == task["title"]
    assert row[12] == task["deliverable"] and row[13] == task["acceptance"]
    assert row[14] == "False" and row[15] == "unfinish"
for row, task in enumerate(tasks, 2):
    assert sheet.cell(row, 4).value == task["title"]
    assert sheet.cell(row, 13).value == task["deliverable"]
    assert sheet.cell(row, 14).value == task["acceptance"]
print("PASS: 40 tasks, 30 API / 6 deferred frontend / 4 QA; acyclic dependencies.")
print("PASS: per-member/day <=6 hours; default calendar 12/10/2026 through 06/11/2026.")
print("PASS: XLSX formulas/flags/owners/PM columns and CSV match the task ledger.")
print("PASS: authorized local Django UI assets exist; original frontend phases and PM gate remain unchanged.")
print("Limit: spreadsheet formulas require Excel/Sheets recalculation; this check does not run backend APIs or prove spreadsheet recalculation/camera accuracy.")
