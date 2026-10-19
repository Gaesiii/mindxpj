"""Convert this project's Vietnamese design Markdown to a formatted Word file."""
from pathlib import Path
import re
import sys
import os

tool_deps = Path(os.environ.get("DOCUMENT_TOOL_DEPS", Path(__file__).parent / "deps"))
sys.path.insert(0, str(tool_deps))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "thiet-ke-he-thong-diem-danh.md"
DESTINATION = SOURCE.with_name("thiet-ke-he-thong-diem-danh-v1.5.docx")


def rich_text(paragraph, text):
    for token in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not token:
            continue
        run = paragraph.add_run(token[2:-2] if token.startswith("**") else
                                token[1:-1] if token.startswith("`") else token)
        if token.startswith("**"):
            run.bold = True
        elif token.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def shade(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def table_from_lines(document, lines):
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    count = len(rows[0])
    if any(len(row) != count for row in rows):
        raise ValueError("Inconsistent Markdown table columns")
    table = document.add_table(rows=1, cols=count)
    table.style = "Table Grid"
    table.autofit = False
    # Reserve most space for descriptive content, keeping IDs compact.
    available_cm = 17.0
    if count == 2:
        widths = [5.0, 12.0]
    elif count == 3:
        widths = [3.0, 7.0, 7.0]
    elif count == 4:
        widths = [2.0, 6.0, 6.0, 3.0]
    else:
        widths = [available_cm / count] * count
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    header = table.rows[0]
    repeat = OxmlElement("w:tblHeader")
    header._tr.get_or_add_trPr().append(repeat)
    for i, value in enumerate(rows[0]):
        cell = header.cells[i]
        shade(cell, "E7EEF7")
        rich_text(cell.paragraphs[0], value)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            rich_text(cells[i].paragraphs[0], value)
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for index, cell in enumerate(row.cells):
            cell.width = Cm(widths[index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.space_before = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build(source=SOURCE, destination=DESTINATION, header_text="THIẾT KẾ ĐIỂM DANH • MVP / PYTHON • BẢN 1.5", version="1.5"):
    content = source.read_text(encoding="utf-8")
    lines = content.splitlines()
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [("Title", 26), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("17365D")
        style.paragraph_format.keep_with_next = True
    document.styles["Heading 1"].paragraph_format.page_break_before = True

    code = document.styles.add_style("Design Code", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Cm(0.25)

    header = section.header.paragraphs[0]
    header.text = header_text
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("61738A")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Trang ")
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)

    document.core_properties.title = lines[0].removeprefix("# ")
    document.core_properties.subject = "Workflow, luồng dữ liệu, kế hoạch module và đặc tả cho agent code"
    document.core_properties.author = "Tài liệu dự án"
    document.core_properties.language = "vi-VN"
    document.core_properties.version = version

    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            in_code = not in_code
            if not in_code:
                document.add_paragraph()
            i += 1
            continue
        if in_code:
            paragraph = document.add_paragraph(line, "Design Code")
            paragraph.paragraph_format.keep_with_next = i + 1 < len(lines) and not lines[i + 1].startswith("```")
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            table_from_lines(document, table_lines)
            continue
        match = re.match(r"^(#{1,4}) (.*)$", line)
        if match:
            level, title = len(match.group(1)), match.group(2)
            if level == 1:
                document.add_paragraph(title, "Title")
            else:
                document.add_heading(title, level=level - 1)
            i += 1
            continue
        if line.startswith("> "):
            paragraph = document.add_paragraph(style="Intense Quote")
            rich_text(paragraph, line[2:])
        elif line.startswith("- [ ] "):
            paragraph = document.add_paragraph(style="List Bullet")
            rich_text(paragraph, "☐ " + line[6:])
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            rich_text(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            rich_text(paragraph, line.strip())
        i += 1
    if in_code:
        raise ValueError("Unclosed code fence")
    document.save(destination)
    print(f"Created: {destination}")
    print(f"Source lines: {len(lines)}; Word tables: {len(document.tables)}; Word paragraphs: {len(document.paragraphs)}")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--destination", type=Path, default=DESTINATION)
    parser.add_argument("--header", default="THIẾT KẾ ĐIỂM DANH • MVP / PYTHON • BẢN 1.5")
    parser.add_argument("--version", default="1.5")
    args = parser.parse_args()
    build(args.source, args.destination, args.header, args.version)
