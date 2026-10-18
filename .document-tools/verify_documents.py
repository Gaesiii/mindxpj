"""Check Word package integrity and content parity with the design source."""
from pathlib import Path
import re
import sys
import os
from zipfile import ZipFile
from xml.etree import ElementTree

tool_deps = Path(os.environ.get("DOCUMENT_TOOL_DEPS", Path(__file__).parent / "deps"))
sys.path.insert(0, str(tool_deps))
from docx import Document

root = Path(__file__).resolve().parents[1]
source = root / "docs/thiet-ke-he-thong-diem-danh.md"
word = source.with_name("thiet-ke-he-thong-diem-danh-v1.4.docx")
markdown = source.read_text(encoding="utf-8")
current_word = source.with_suffix(".docx")
archived_source = root / "docs/archive/thiet-ke-he-thong-diem-danh-v1.2.md"
archived_word = source.with_name("thiet-ke-he-thong-diem-danh-v1.2.docx")
assert archived_source.is_file() and archived_word.is_file(), "Missing v1.2 historical artifact"
assert "**Phiên bản:** 1.2" in archived_source.read_text(encoding="utf-8"), "Archive is not v1.2"
v13_source = root / "docs/archive/thiet-ke-he-thong-diem-danh-v1.3.md"
v13_word = root / "docs/archive/thiet-ke-he-thong-diem-danh-v1.3.docx"
assert v13_source.is_file() and v13_word.is_file(), "Missing v1.3 archive"
assert "**Phiên bản 1.3" in v13_source.read_text(encoding="utf-8")
assert (source.with_name("thiet-ke-he-thong-diem-danh-v1.3.docx")).read_bytes() == v13_word.read_bytes(), "v1.3 Word archive changed"
assert current_word.read_bytes() == word.read_bytes(), "Current Word differs from v1.4 Word"


def validate_package(path):
    with ZipFile(path) as package:
        assert package.testzip() is None, f"Corrupt Word ZIP package: {path.name}"
        for name in package.namelist():
            if name.endswith(".xml"):
                ElementTree.fromstring(package.read(name))
        return ElementTree.fromstring(package.read("word/document.xml"))


document_xml = validate_package(word)
validate_package(current_word)
validate_package(archived_word)
validate_package(v13_word)

ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
# XML text below excludes formatting metadata.
paragraphs = ["".join(t.text or "" for t in p.findall(".//w:t", ns))
              for p in document_xml.findall(".//w:p", ns)]
word_text = "\n".join(paragraphs)
assert "�" not in markdown + word_text, "Broken Unicode"
assert markdown.count("```") % 2 == 0, "Unbalanced code fences"
assert not re.search(r"\b(TODO|TBD|FIXME)\b", markdown), "Unfinished placeholder"

expected = []
in_code = False
for line in markdown.splitlines():
    if line.startswith("```"):
        in_code = not in_code
        continue
    if not line.strip():
        continue
    if in_code:
        expected.append(line)
        continue
    if line.startswith("|"):
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        candidates = cells
    else:
        line = re.sub(r"^#{1,4} ", "", line)
        line = re.sub(r"^> ", "", line)
        line = re.sub(r"^- \[ \] ", "☐ ", line)
        line = re.sub(r"^- ", "", line)
        candidates = [line.strip()]
    for value in candidates:
        value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
        value = re.sub(r"`([^`]+)`", r"\1", value)
        expected.append(value)

missing = [value for value in expected if value not in word_text]
assert not missing, f"Word missing content: {missing[:5]}"
sections = re.findall(r"^## (\d+)\. ", markdown, re.M)
assert sections == [str(n) for n in range(1, 6)], sections
expected_headings = re.findall(r"^## (.+)$", markdown, re.M)
assert len(expected_headings) == 5
assert "**Phiên bản 1.4" in markdown
for required in ("Python 3.12/Django 5.2", "OpenCV", "YuNet", "SFace", "HSV", "VISION", "QR_ENABLED=false",
                 "NOT_APPROVED", "112x112", "512byte", "120 giây", "identity_confirmed=true", "manual",
                 "start_at <= now <= close_at"):
    assert required in markdown, f"Missing v1.4 scope constraint: {required}"
document = Document(word)
expected_tables = len(re.findall(r"^\|(?:\s*:?-{3,}:?\s*\|)+\s*$", markdown, re.M))
assert len(document.tables) == expected_tables
assert expected_tables == 1, "v1.4 ownership table missing or unexpectedly changed"
assert "UI demo cục bộ đã triển khai" in markdown and "làm luôn UI để demo" in markdown
assert "không phải phê duyệt PM API" in markdown and "demo_ui" in markdown
assert not re.search(r"TypeScript|NestJS|React|npm run|\.tsx?\b|\.spec\.ts", markdown), "Old stack remains"
actual_headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
assert actual_headings == expected_headings, (actual_headings, expected_headings)
print("PASS: Word ZIP and XML are valid; Vietnamese Unicode preserved.")
print(f"PASS: {len(expected)} source content fragments present in Word.")
print(f"PASS: v1.4 scope and {len(sections)} numbered sections match the current source.")
print("PASS: v1.2/v1.3 Markdown/Word retained; current Word is byte-identical to v1.4 Word.")
print(f"PASS: {len(document.tables)} tables; all section headings preserved.")
print("Limit: checks verify content and document structure, not visual page rendering in Microsoft Word.")
