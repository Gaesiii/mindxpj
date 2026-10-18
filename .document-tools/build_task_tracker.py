"""Render the task JSON as a portable XLSX tracker and raw CSV."""
from pathlib import Path
from datetime import datetime
import csv
import json
import sys

sys.path.insert(0, str(Path(__file__).parent / "deps"))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/project-management"
DATA = json.loads((OUT / "tasks.json").read_text(encoding="utf-8"))
TASKS = DATA["tasks"]
HEADERS = ["ID", "Module", "Phase", "Task", "Owner", "Reviewer", "Start day",
           "End day", "Start date", "Due date", "Hours", "Dependencies", "Deliverable",
           "Acceptance", "Finish?", "Status", "Ready / Blocked", "PM Comment",
           "Evidence / PR", "Actual Finish", "Dev Note"]


def ready_formula(task, row, end):
    conditions = [f'COUNTIFS($A$2:$A${end},"{dep}",$P$2:$P${end},"finish")=1'
                  for dep in task["deps"]]
    if task["phase"] != "API":
        conditions += ['Config!$B$6="APPROVED"',
                       f'COUNTIFS($C$2:$C${end},"API",$P$2:$P${end},"unfinish")=0',
                       f'COUNTIF($C$2:$C${end},"API")>0']
    ready = "AND(" + ",".join(conditions) + ")" if conditions else "TRUE"
    return f'=IF(OR(O{row}=TRUE,O{row}="TRUE"),"done",IF({ready},"ready","blocked"))'


def style(sheet, widths):
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 34
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build():
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    config = wb.active
    config.title = "Config"
    config.append(["Setting", "Value", "Notes"])
    rows = [
        ["Start date", datetime.fromisoformat(DATA["start_date"]), "Đổi ngày bắt đầu; nghỉ cuối tuần"],
        ["Member 1", DATA["members"][0], "Đổi tên tại đây"],
        ["Member 2", DATA["members"][1], "Đổi tên tại đây"],
        ["PM", DATA["pm"], "Người review/ghi comment"],
        ["API_Gate", "NOT_APPROVED", "PM duyệt sau khi API có evidence"],
        ["Capacity hours/day", DATA["hours_per_member_per_day"], "Giả định mỗi người"],
        ["Working days", DATA["working_days"], "Khung một tháng"],
        ["Scope", "MVP only", "v1.3; local UI user-authorized before gate; PM NOT_APPROVED"],
    ]
    for row in rows:
        config.append(row)
    config["B2"].number_format = "dd/mm/yyyy"
    config["D1"] = "Holidays (optional)"
    for row in range(2, 31):
        config.cell(row, 4).number_format = "dd/mm/yyyy"
    gate = DataValidation(type="list", formula1='"NOT_APPROVED,APPROVED"', allow_blank=False)
    gate.errorTitle, gate.error, gate.showErrorMessage = "API gate", "Chọn giá trị trong danh sách", True
    config.add_data_validation(gate)
    gate.add(config["B6"])
    style(config, {"A": 25, "B": 25, "C": 55, "D": 25})

    sheet = wb.create_sheet("Tasks")
    sheet.append(HEADERS)
    end = len(TASKS) + 1
    for r, task in enumerate(TASKS, 2):
        reviewer = '=Config!$B$5' if task["start_day"] >= 15 and task["phase"] != "FRONTEND_LATER" else f'=Config!$B${4 if task["owner"] == 1 else 3}'
        sheet.append([
            task["id"], task["module"], task["phase"], task["title"],
            f'=Config!$B${3 if task["owner"] == 1 else 4}', reviewer,
            task["start_day"], task["end_day"],
            f'=WORKDAY(Config!$B$2-1,G{r},Config!$D$2:$D$30)',
            f'=WORKDAY(Config!$B$2-1,H{r},Config!$D$2:$D$30)',
            task["hours"], ", ".join(task["deps"]), task["deliverable"], task["acceptance"],
            False, f'=IF(OR(O{r}=TRUE,O{r}="TRUE"),"finish","unfinish")',
            ready_formula(task, r, end), "", "", "", "",
        ])
        sheet.row_dimensions[r].height = 80
        for col in (9, 10):
            sheet.cell(r, col).number_format = "dd/mm/yyyy"
        sheet.cell(r, 20).number_format = "dd/mm/yyyy hh:mm"
    widths = {"A": 9, "B": 24, "C": 23, "D": 60, "E": 16, "F": 16, "G": 12,
              "H": 12, "I": 16, "J": 16, "K": 10, "L": 24, "M": 65, "N": 75,
              "O": 14, "P": 15, "Q": 18, "R": 50, "S": 40, "T": 24, "U": 45}
    style(sheet, widths)
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = f"A1:U{end}"
    checked = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
    checked.showErrorMessage = True
    checked.error = "Chọn TRUE khi hoàn tất, FALSE khi chưa hoàn tất"
    sheet.add_data_validation(checked)
    checked.add(f"O2:O{end}")
    sheet.conditional_formatting.add(f"A2:U{end}", FormulaRule(formula=['$P2="finish"'], fill=PatternFill("solid", fgColor="DCFCE7")))
    sheet.conditional_formatting.add(f"A2:U{end}", FormulaRule(formula=['AND($P2="unfinish",$J2<TODAY())'], fill=PatternFill("solid", fgColor="FEE2E2")))
    sheet.conditional_formatting.add(f"Q2:Q{end}", FormulaRule(formula=['$Q2="blocked"'], fill=PatternFill("solid", fgColor="FFF7ED")))

    dashboard = wb.create_sheet("Dashboard")
    dashboard.append(["Metric", "Value"])
    metrics = [
        ["Total tasks", f'=COUNTA(Tasks!A2:A{end})'],
        ["Finish", f'=COUNTIF(Tasks!P2:P{end},"finish")'],
        ["Unfinish", f'=COUNTIF(Tasks!P2:P{end},"unfinish")'],
        ["Overall progress", '=IFERROR(B3/B2,0)'],
        ["API tasks", f'=COUNTIF(Tasks!C2:C{end},"API")'],
        ["API finish", f'=COUNTIFS(Tasks!C2:C{end},"API",Tasks!P2:P{end},"finish")'],
        ["API progress", '=IFERROR(B7/B6,0)'],
        ["Frontend planned", f'=COUNTIF(Tasks!C2:C{end},"FRONTEND_LATER")'],
        ["API_Gate", '=Config!B6'],
        ["Frontend unlocked", '=IF(AND(B7=B6,B6>0,Config!B6="APPROVED"),"YES","NO")'],
        ["Owner", "Progress"],
        ['=Config!B3', f'=IFERROR(COUNTIFS(Tasks!E2:E{end},A13,Tasks!P2:P{end},"finish")/COUNTIF(Tasks!E2:E{end},A13),0)'],
        ['=Config!B4', f'=IFERROR(COUNTIFS(Tasks!E2:E{end},A14,Tasks!P2:P{end},"finish")/COUNTIF(Tasks!E2:E{end},A14),0)'],
        ["mem1 planned hours", f'=SUMIF(Tasks!E2:E{end},Config!B3,Tasks!K2:K{end})'],
        ["mem2 planned hours", f'=SUMIF(Tasks!E2:E{end},Config!B4,Tasks!K2:K{end})'],
        ["Note", "Tick finish after tests/review/evidence; PM decides API gate."],
    ]
    for row in metrics:
        dashboard.append(row)
    for coordinate in ("B5", "B8", "B13", "B14"):
        dashboard[coordinate].number_format = "0.0%"
    style(dashboard, {"A": 32, "B": 65})

    modules = wb.create_sheet("Modules")
    modules.append(["Module", "Owner", "Scope"])
    for module, owner, scope in [
        ("Accounts", 1, "Auth/session/role/scope"),
        ("Academics", 1, "Grade/class/one homeroom teacher"),
        ("Students", 1, "Profile/class/private photo"),
        ("Reports", 1, "Scoped summaries/XLSX/CSV"),
        ("Schedules", 2, "Weekly rules/session/roster snapshot"),
        ("Attendance / Uniform", 2, "Arrival/manual/verification/uniform/close"),
        ("Vision / QR legacy", 2, "OpenCV staff confirm; QR disabled by default"),
    ]:
        modules.append([module, f'=Config!B{3 if owner == 1 else 4}', scope])
    style(modules, {"A": 28, "B": 18, "C": 70})
    wb.save(OUT / "task-tracker.xlsx")

    with (OUT / "tasks-import.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADERS)
        for t in TASKS:
            reviewer = DATA["pm"] if t["start_day"] >= 15 and t["phase"] != "FRONTEND_LATER" else DATA["members"][1 if t["owner"] == 1 else 0]
            writer.writerow([t["id"], t["module"], t["phase"], t["title"],
                             DATA["members"][t["owner"]-1], reviewer, t["start_day"], t["end_day"],
                             "", "", t["hours"], ", ".join(t["deps"]), t["deliverable"], t["acceptance"],
                             False, "unfinish", "", "", "", "", ""])
    print(f"Created task-tracker.xlsx and tasks-import.csv: {len(TASKS)} tasks.")


if __name__ == "__main__":
    build()
