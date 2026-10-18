"""Validate the student guide's Word package, source parity and teaching structure."""
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile
import json
import re
import sys

sys.path.insert(0, str(Path(__file__).parent / "deps"))
from docx import Document

root = Path(__file__).resolve().parents[1]
source = root / "docs/huong-dan-workflow-module.md"
word = source.with_suffix(".docx")
markdown = source.read_text(encoding="utf-8")
with ZipFile(word) as package:
    assert package.testzip() is None
    for name in package.namelist():
        if name.endswith(".xml"):
            ElementTree.fromstring(package.read(name))
    xml = ElementTree.fromstring(package.read("word/document.xml"))
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
text = "\n".join("".join(t.text or "" for t in p.findall(".//w:t", ns))
                 for p in xml.findall(".//w:p", ns))
assert "�" not in markdown + text
assert not re.search(r"học viên|học sinh lớp 10\b|dành cho.*lớp 10\b|nền tảng.*còn mới", markdown + text, re.I)
assert markdown.count("```") % 2 == 0
assert not re.search(r"\b(TODO|TBD|FIXME)\b", markdown)
expected = []
in_code = False
for line in markdown.splitlines():
    if line.startswith("```"):
        in_code = not in_code
        continue
    if not line.strip():
        continue
    if in_code:
        expected.append(line)
        continue
    if line.startswith("|"):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        values = cells
    else:
        line = re.sub(r"^#{1,4} ", "", line)
        line = re.sub(r"^> ", "", line)
        line = re.sub(r"^- \[ \] ", "☐ ", line)
        line = re.sub(r"^- ", "", line)
        values = [line.strip()]
    for value in values:
        value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
        value = re.sub(r"`([^`]+)`", r"\1", value)
        expected.append(value)
missing = [part for part in expected if part not in text]
assert not missing, missing[:3]
headings = re.findall(r"^## (\d+)\.", markdown, re.M)
assert headings == [str(i) for i in range(1, 15)]
modules = re.findall(r"^## \d+\. (M\d{2})", markdown, re.M)
assert modules == [f"M{i:02d}" for i in range(1, 9)]
for index, module in enumerate(modules):
    start = re.search(rf"^## \d+\. {module}", markdown, re.M).start()
    end = re.search(r"^## \d+\.", markdown[start+1:], re.M)
    section = markdown[start:start+1+end.start()] if end else markdown[start:]
    assert "Workflow" in section, module
    assert "Mẫu yêu cầu agent" in section, module
    assert "Task:" in section, module
    assert "kiểm" in section.lower(), module
for keyword in ["Authentication", "Authorization", "Foreign key", "Migration", "ORM",
                "Transaction", "Idempotency", "CSRF", "Version", "snapshot",
                "Integration", "API gate", "Virtual environment"]:
    assert keyword in markdown, keyword
examples = re.findall(r"```json\n(.*?)\n```", markdown, re.S)
for example in examples:
    json.loads(example)
doc = Document(word)
assert len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) == 14
assert len(doc.tables) == len(re.findall(r"^\|(?:\s*:?-{3,}:?\s*\|)+\s*$", markdown, re.M))
print(f"PASS: Word ZIP/XML and Vietnamese text; {len(expected)} source fragments preserved.")
print(f"PASS: 14 chapters, 8 modules with workflows/tasks/checks/agent prompts; {len(doc.tables)} tables.")
print(f"PASS: key terminology explained; {len(examples)} JSON examples parse.")
print("Limit: checks validate content and structure, not visual page rendering or application APIs.")
