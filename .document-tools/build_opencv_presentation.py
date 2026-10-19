"""Build the Vietnamese 5-7 minute OpenCV presentation, with speaker notes.

Tool-only dependency setup (Python 3.12; not a backend dependency):
backend/.venv/Scripts/python.exe -m pip install --target .document-tools/presentation-deps python-pptx==1.0.2
backend/.venv/Scripts/python.exe .document-tools/build_opencv_presentation.py
"""
from pathlib import Path
import sys
import os

ROOT = Path(__file__).resolve().parents[1]
presentation_deps = Path(os.environ.get('PRESENTATION_TOOL_DEPS', Path(__file__).parent / 'presentation-deps'))
sys.path.insert(0, str(presentation_deps))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = ROOT / 'docs/thuyet-trinh-opencv.pptx'
NAVY = '102338'
WHITE = 'F5F8FC'
MUTED = 'B6C8DB'
TEAL = '55D5C6'
CARD = '1C354D'
AMBER = 'FFD18A'

import json
SLIDES = json.loads((ROOT / '.document-tools/presentation_v15.json').read_text(encoding='utf-8'))


def color(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    return shape


def text(slide, x, y, w, h, value, size=21, fill=WHITE, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(fill)
        p.space_after = Pt(14)
    return shape


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for number, (title, subtitle, cards, notes) in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color(NAVY)
        text(slide, .65, .48, 12, .7, title, 33, bold=True)
        text(slide, .65, 1.25, 12, .65, subtitle, 19, TEAL)
        if len(cards) == 4:
            for i, (heading, lines) in enumerate(cards):
                x = .65 + i * 3.12
                box(slide, x, 2.2, 2.95, 3.7, CARD)
                text(slide, x+.18, 2.42, 2.6, .55, heading, 22, TEAL, True)
                text(slide, x+.18, 3.12, 2.6, 2.6, '\n'.join(lines), 20)
        else:
            for i, (heading, lines) in enumerate(cards):
                x = .65 + i * 6.17
                box(slide, x, 2.12, 5.99, 4.37, CARD)
                text(slide, x+.25, 2.35, 5.45, .55, heading, 24, TEAL, True)
                text(slide, x+.25, 3.12, 5.43, 3.08, '\n'.join('• '+line for line in lines), 21)
        text(slide, .65, 6.98, 11.5, .25, 'OpenCV MVP v1.5 | Tự điểm danh | Giáo viên xử lý ngoại lệ', 11, MUTED)
        text(slide, 12.05, 6.96, .65, .3, f'{number}/7', 12, AMBER)
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(OUT)
    # Reopen and confirm every authored title, card and note survived serialization.
    reopened = Presentation(OUT)
    assert len(reopened.slides) == 7
    for slide, (title, subtitle, cards, notes) in zip(reopened.slides, SLIDES):
        all_text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert title in all_text and subtitle in all_text
        for heading, lines in cards:
            assert heading in all_text
            assert all(line in all_text for line in lines)
        assert notes in slide.notes_slide.notes_text_frame.text
    print('PASS: created and reopened 7 slides; all titles, card text and speaker notes match.')
    print('Limit: no PowerPoint rendering or live presentation/camera test performed.')
    print(str(OUT))


if __name__ == '__main__':
    build()
