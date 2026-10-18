from django.test import SimpleTestCase

from school_app.reports.services import summarize


class ReportAggregateTests(SimpleTestCase):
    def test_forty_roster_example(self):
        rows = [dict(status='PRESENT', arrival_status='ON_TIME' if i < 30 else 'LATE',
                     uniform_status='COMPLIANT' if i < 27 else 'NON_COMPLIANT' if i < 31 else 'NOT_CHECKED')
                for i in range(34)]
        rows += [dict(status=status, arrival_status='NOT_APPLICABLE', uniform_status='NOT_APPLICABLE')
                 for status in ['ABSENT'] * 3 + ['UNMARKED'] * 3]
        result = summarize(rows)
        self.assertEqual({k: result[k] for k in ('expected', 'present', 'absent', 'unmarked', 'on_time', 'late',
                         'compliant', 'non_compliant', 'not_checked', 'checked')},
                         dict(expected=40, present=34, absent=3, unmarked=3, on_time=30, late=4,
                              compliant=27, non_compliant=4, not_checked=3, checked=31))
        self.assertEqual(result['presence_rate'], 34/40)
        self.assertEqual(result['late_rate'], 4/34)
        self.assertEqual(result['violation_rate'], 4/31)
        self.assertEqual(result['check_coverage'], 31/34)

    def test_zero_denominators_are_null(self):
        result = summarize([])
        for key in ('presence_rate', 'late_rate', 'violation_rate', 'check_coverage'):
            self.assertIsNone(result[key])

from datetime import date, datetime, timedelta, timezone
from io import BytesIO, StringIO
import csv
import uuid
from django.test import TestCase, override_settings
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from school_app.academics.models import Grade, Classroom
from school_app.students.models import Student
from school_app.schedules.models import Session, SessionRoster
from school_app.attendance.models import Attendance
from school_app.qr.models import QrChallenge, QrScan


@override_settings(ROOT_URLCONF='school_app.reports.urls')
class ReportApiTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.teacher = User.objects.create_user(username='reportteacher', role='HOMEROOM_TEACHER')
        self.other = User.objects.create_user(username='reportother', role='HOMEROOM_TEACHER')
        self.supervisor = User.objects.create_user(username='reportsuper', role='SUPERVISOR')
        self.pupil = User.objects.create_user(username='reportpupil', role='STUDENT')
        self.grade = Grade.objects.create(code='10', name='Khối 10')
        self.other_grade = Grade.objects.create(code='11', name='Khối 11')
        self.classroom = Classroom.objects.create(code='10A', name='Lớp 10A', grade=self.grade, teacher=self.teacher)
        self.other_class = Classroom.objects.create(code='11B', name='Lớp 11B', grade=self.other_grade, teacher=self.other)
        self.cutoff = datetime(2026, 10, 12, 0, 5, tzinfo=timezone.utc)
        self.session = self.new_session(self.classroom)
        self.rows = [self.new_row(self.session, i) for i in range(3)]
        Attendance.objects.filter(roster=self.rows[0]).update(status='PRESENT', source='MANUAL',
                 arrived_at=self.cutoff, uniform_status='COMPLIANT', observed_uniform='SHIRT')
        Attendance.objects.filter(roster=self.rows[1]).update(status='ABSENT')
        outside = self.new_session(self.other_class)
        self.new_row(outside, 10)
        Session.objects.create(classroom=self.classroom, school_date=date(2026,10,13), slot='MORNING')
        self.client.force_login(self.teacher)
        self.filters = dict(date_from='2026-10-12', date_to='2026-10-13')

    def new_session(self, classroom, **kwargs):
        return Session.objects.create(classroom=classroom, school_date=date(2026,10,12), slot='MORNING',
            status='OPEN', start_at=self.cutoff-timedelta(minutes=5), late_cutoff_at=self.cutoff,
            close_at=self.cutoff+timedelta(hours=1), required_uniform_snapshot='SHIRT', **kwargs)

    def new_row(self, session, i):
        student = Student.objects.create(student_code=f'S{i}', full_name=f'Nguyễn An {i}', classroom=session.classroom)
        row = SessionRoster.objects.create(session=session, student=student, student_code_snapshot=student.student_code,
            full_name_snapshot=student.full_name, class_code_snapshot=session.classroom.code,
            class_name_snapshot=session.classroom.name, grade_code_snapshot=session.classroom.grade.code,
            grade_name_snapshot=session.classroom.grade.name)
        Attendance.objects.create(roster=row)
        return row

    def get(self, endpoint='attendance', **filters):
        return self.client.get('/reports/'+endpoint, {**self.filters, **filters})

    def test_scope_snapshots_cutoff_and_filters(self):
        self.rows[0].student.full_name = 'Changed'
        self.rows[0].student.classroom = self.other_class
        self.rows[0].student.save()
        result = self.get().json()['items']
        self.assertEqual(len(result), 3)
        first = next(row for row in result if row['roster_id'] == str(self.rows[0].id))
        self.assertEqual(first['full_name_snapshot'], 'Nguyễn An 0')
        self.assertEqual(first['class_name_snapshot'], 'Lớp 10A')
        self.assertEqual(first['arrival_status'], 'ON_TIME')
        self.assertEqual(first['required_uniform_snapshot'], 'SHIRT')
        for criteria in (dict(status='PRESENT'), dict(grade_id=str(self.grade.id), slot='MORNING', status='PRESENT')):
            self.assertEqual(len(self.get(**criteria).json()['items']), 1)
        self.assertEqual(self.get(slot='AFTERNOON').json()['items'], [])
        self.assertEqual(self.get('summary', status='PRESENT').json()['expected'], 1)

    def test_roles_reassignment_and_explicit_outside_filters(self):
        for endpoint in ('attendance', 'summary', 'export'):
            extras = dict(format='csv') if endpoint == 'export' else {}
            for field, value in [('class_id', self.other_class.id), ('grade_id', self.other_grade.id),
                                 ('class_id', uuid.uuid4()), ('grade_id', uuid.uuid4())]:
                self.assertEqual(self.get(endpoint, **extras, **{field:str(value)}).status_code, 403)
            self.client.force_login(self.pupil)
            self.assertEqual(self.get(endpoint, **extras).status_code, 403)
            self.client.logout()
            self.assertEqual(self.get(endpoint, **extras).status_code, 401)
            self.client.force_login(self.teacher)
        self.classroom.teacher = self.other
        self.classroom.save()
        self.assertEqual(self.get().json()['items'], [])
        self.client.force_login(self.other)
        self.assertEqual(len(self.get().json()['items']), 4)
        self.client.force_login(self.supervisor)
        self.assertEqual(len(self.get().json()['items']), 4)
        self.supervisor.is_active = False
        self.supervisor.save()
        self.assertEqual(self.get().status_code, 401)

    def test_invalid_filters(self):
        cases = [dict(date_from=''), dict(date_to=''), dict(date_from='2026-10-14'),
                 dict(date_from='12/10/2026'), dict(date_from='2026-1-2'), dict(status='PENDING'),
                 dict(status='OPEN'), dict(slot='NIGHT'), dict(class_id='bad'), dict(grade_id='bad'),
                 dict(student_id=str(uuid.uuid4())), dict(format='csv')]
        for criteria in cases:
            with self.subTest(criteria=criteria):
                self.assertEqual(self.get(**criteria).status_code, 422)
        self.assertEqual(self.client.get('/reports/attendance').status_code, 422)
        self.assertEqual(self.get('export', format='pdf').status_code, 422)
        self.assertEqual(self.get('export').status_code, 422)
        self.assertEqual(self.client.get('/reports/attendance', [('date_from','2026-10-12'),
            ('date_from','2026-10-13'), ('date_to','2026-10-13')]).status_code, 422)

    def test_pending_not_present_and_only_relevant_confirmed_scan(self):
        challenge = QrChallenge.objects.create(session=self.session, token_hash='report-hash', issued_by=self.teacher,
            issued_at=self.cutoff, expires_at=self.cutoff+timedelta(seconds=60))
        for i, status in [(0,'CONFIRMED'), (0,'REJECTED'), (2,'PENDING')]:
            QrScan.objects.create(roster=self.rows[i], challenge=challenge, student_user=self.pupil,
                scanned_at=self.cutoff + timedelta(seconds=10 if status == 'REJECTED' else 0),
                expires_at=self.cutoff+timedelta(minutes=15), status=status)
        Attendance.objects.filter(roster=self.rows[0]).update(source='QR')
        rows = self.get().json()['items']
        first = next(row for row in rows if row['roster_id'] == str(self.rows[0].id))
        self.assertEqual(first['scanned_at'], self.cutoff.isoformat())
        self.assertEqual(self.get('summary').json()['present'], 1)
        Attendance.objects.filter(roster=self.rows[0]).update(source='MANUAL')
        first = next(row for row in self.get().json()['items'] if row['roster_id'] == str(self.rows[0].id))
        self.assertIsNone(first['scanned_at'])

    def test_exports_unicode_formula_defense_metadata_and_reconcile(self):
        risky = ['=1+1', '+cmd', '-cmd', '@cmd', ' \t=1+1', '\n@cmd', '\x01=cmd', '\u2003\x01=cmd', '\u200b=cmd']
        for value in risky:
            SessionRoster.objects.filter(pk=self.rows[0].pk).update(full_name_snapshot=value)
            response = self.get('export', format='xlsx')
            self.assertEqual(response.status_code, 200)
            book = load_workbook(BytesIO(response.content))
            self.assertEqual(book.sheetnames, ['Chi_tiet','Tong_hop','Thong_tin'])
            sheet = book['Chi_tiet']
            headers = [cell.value for cell in sheet[1]]
            cell = sheet.cell(2, headers.index('Họ tên')+1)
            # Stable order puts student S0 first.
            self.assertEqual(cell.data_type, 's')
            self.assertTrue(cell.value.startswith("'"))
            response = self.get('export', format='csv')
            self.assertTrue(response.content.startswith(b'\xef\xbb\xbf'))
            records = list(csv.reader(StringIO(response.content.decode('utf-8-sig'))))
            self.assertEqual(records[1][headers.index('Họ tên')][0], "'")
        info = list(book['Thong_tin'].values)
        self.assertTrue(any('Asia/Ho_Chi_Minh' in str(row) for row in info))
        self.assertTrue(any('2026-10-12' in str(row) for row in info))
        summary = dict(book['Tong_hop'].values)
        self.assertEqual(summary['present'], self.get('summary').json()['present'])
        csv_response = self.get('export', format='csv')
        text = csv_response.content.decode('utf-8-sig')
        self.assertIn('Khối 10', text)
        self.assertIn('07:05:00+07:00', text)
        for private in ('password', 'photo_path', 'token_hash', 'report-hash'):
            self.assertNotIn(private, text)
        self.assertEqual(csv_response['Content-Disposition'], 'attachment; filename="attendance.csv"')

    def test_export_limit_on_real_roster_count(self):
        students = [Student(student_code=f'LIMIT{i}', full_name='Limit', classroom=self.classroom) for i in range(4998)]
        Student.objects.bulk_create(students)
        rows = [SessionRoster(session=self.session, student=s, student_code_snapshot=s.student_code,
            full_name_snapshot='Limit', class_code_snapshot='10A', class_name_snapshot='Lớp 10A',
            grade_code_snapshot='10', grade_name_snapshot='Khối 10') for s in students]
        SessionRoster.objects.bulk_create(rows)
        Attendance.objects.bulk_create([Attendance(roster=r) for r in rows])
        response = self.get('export', format='csv')
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()['code'], 'EXPORT_LIMIT_EXCEEDED')
        self.assertIn('chia', response.json()['message'])
        self.assertEqual(len(self.get().json()['items']), 5001)
        Attendance.objects.filter(roster=rows[-1]).delete()
        rows[-1].delete()
        allowed = self.get('export', format='csv')
        self.assertEqual(allowed.status_code, 200)
        self.assertEqual(len(list(csv.reader(StringIO(allowed.content.decode('utf-8-sig'))))), 5001)

    def test_xml_invalid_noncharacters_are_sanitized(self):
        SessionRoster.objects.filter(pk=self.rows[0].pk).update(full_name_snapshot='Nguyễn\ufffe\uffff An')
        response = self.get('export', format='xlsx')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book['Chi_tiet'].cell(2, 4).value, 'Nguyễn An')
