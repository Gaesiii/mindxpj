from django.test import override_settings
"""Real API workflow reconciliation: no report fixtures substitute for writes."""
import csv
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from unittest.mock import patch
from django.test import TestCase
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from school_app.academics.models import Grade, Classroom
from school_app.students.models import Student
from school_app.schedules.models import ScheduleRule, SessionRoster


@override_settings(QR_ENABLED=True)
class ReportReconciliationTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.teacher = User.objects.create_user(username='reconcileteacher', role='HOMEROOM_TEACHER')
        grade = Grade.objects.create(code='R10', name='Khối mười')
        self.classroom = Classroom.objects.create(code='R10A', name='Lớp mười A', grade=grade, teacher=self.teacher)
        self.pupils = []
        for i in range(4):
            student = Student.objects.create(student_code=f'R{i}', full_name=f'Học sinh {i}', classroom=self.classroom)
            self.pupils.append(User.objects.create_user(username=f'reconcilepupil{i}', role='STUDENT', student=student))
        ScheduleRule.objects.create(classroom=self.classroom, weekday=0, slot='MORNING', start_time='07:00',
            close_time='08:00', grace_minutes=5, required_uniform='SHIRT')
        self.now = datetime(2026,10,12,0,4,tzinfo=timezone.utc)
        self.clock = patch('django.utils.timezone.now', return_value=self.now)
        self.mock_now = self.clock.start()
        self.addCleanup(self.clock.stop)
        self.client.force_login(self.teacher)
        response = self.write('post', '/sessions', {'class_id':str(self.classroom.id),
                              'school_date':'2026-10-12','slot':'MORNING'})
        self.session_id = response['id']
        self.write('post', f'/sessions/{self.session_id}/open', {'expected_version':0})
        self.rosters = list(SessionRoster.objects.filter(session_id=self.session_id).order_by('student_code_snapshot'))
        self.filters = {'date_from':'2026-10-12','date_to':'2026-10-12'}

    def write(self, method, path, body, key=None):
        kwargs = {'content_type':'application/json'}
        if key:
            kwargs['HTTP_IDEMPOTENCY_KEY'] = key
        response = getattr(self.client,method)('/api/v1'+path, body, **kwargs)
        self.assertIn(response.status_code, (200,201), response.content)
        return response.json()

    def report(self, endpoint):
        response = self.client.get('/api/v1/reports/'+endpoint, self.filters)
        self.assertEqual(response.status_code,200,response.content)
        return response.json()

    def test_qr_manual_uniform_close_details_summary_and_both_exports_reconcile(self):
        challenge = self.write('post', f'/sessions/{self.session_id}/qr-challenges', {})
        token = challenge['token']
        self.client.force_login(self.pupils[0])
        scan = self.write('post', '/qr-scans', {'token':token}, 'reconcile-scan-0')
        self.client.force_login(self.pupils[3])
        self.write('post', '/qr-scans', {'token':token}, 'reconcile-scan-3')
        self.client.force_login(self.teacher)
        before = self.report('summary')
        self.assertEqual((before['expected'],before['present'],before['unmarked']),(4,0,4))
        self.mock_now.return_value = self.now + timedelta(minutes=4)
        self.write('post', f"/qr-scans/{scan['id']}/confirm", {'expected_attendance_version':0,
            'expected_scan_version':0,'identity_confirmed':True,
            'uniform':{'observed_uniform':'SHIRT','uniform_status':'COMPLIANT'}}, 'reconcile-confirm')
        self.write('patch', f'/attendance/{self.rosters[1].attendance.id}', {'expected_version':0,
            'status':'PRESENT','arrived_at':(self.now+timedelta(minutes=2)).isoformat(),
            'reason':'Ghi thủ công khi không có điện thoại'}, 'reconcile-manual')
        self.write('patch', f'/attendance/{self.rosters[1].attendance.id}/uniform', {'expected_version':1,
            'observed_uniform':'SPORTS','uniform_status':'NON_COMPLIANT','reason':'Yêu cầu sơ mi'}, 'reconcile-uniform')
        self.write('patch', f'/attendance/{self.rosters[2].attendance.id}', {'expected_version':0,
            'status':'PRESENT','arrived_at':self.now.isoformat(),'reason':'Nhập hồi tố'}, 'reconcile-notchecked')
        opened = self.report('summary')
        self.assertEqual({key:opened[key] for key in ('expected','present','absent','unmarked','on_time','late',
            'compliant','non_compliant','not_checked','checked')}, dict(expected=4,present=3,absent=0,
            unmarked=1,on_time=2,late=1,compliant=1,non_compliant=1,not_checked=1,checked=2))
        details = self.report('attendance')['items']
        qr_row = next(row for row in details if row['roster_id']==str(self.rosters[0].id))
        self.assertEqual(datetime.fromisoformat(qr_row['arrived_at']),self.now)
        self.assertEqual(qr_row['scanned_at'],qr_row['arrived_at'])
        self.assertEqual(qr_row['arrival_status'],'ON_TIME')
        self.assertNotEqual(qr_row['verified_at'],qr_row['arrived_at'])
        closed = self.write('post', f'/sessions/{self.session_id}/close', {'expected_version':1})
        self.assertEqual(closed['converted_unmarked'],1)
        self.assertEqual(closed['expired_pending'],1)
        summary = self.report('summary')
        self.assertEqual(summary['present'],3)
        self.assertEqual(summary['absent'],1)
        self.assertEqual(summary['unmarked'],0)
        self.assertEqual(summary['not_checked'],1)
        self.assertEqual(summary['violation_rate'],0.5)
        self.assertEqual(summary['presence_rate'],0.75)
        result = self.client.get('/api/v1/reports/export',{**self.filters,'format':'xlsx'})
        self.assertEqual(result.status_code,200)
        book=load_workbook(BytesIO(result.content))
        self.assertEqual(dict(book['Tong_hop'].values),summary)
        self.assertEqual(book['Chi_tiet'].max_row,5)
        result = self.client.get('/api/v1/reports/export',{**self.filters,'format':'csv'})
        self.assertEqual(result.status_code,200)
        rows=list(csv.DictReader(StringIO(result.content.decode('utf-8-sig'))))
        self.assertEqual(len(rows),4)
        self.assertEqual(sum(row['Trạng thái']=='PRESENT' for row in rows),summary['present'])
        self.assertEqual(sum(row['Kết quả đồng phục']=='NOT_CHECKED' for row in rows),1)
