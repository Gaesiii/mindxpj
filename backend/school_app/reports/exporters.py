"""Synchronous small exports with explicit spreadsheet injection protection."""
import csv
import json
import unicodedata
from datetime import datetime
from io import BytesIO, StringIO
from zoneinfo import ZoneInfo
from django.http import HttpResponse
from openpyxl import Workbook

COLUMNS = (
    ('school_date', 'Ngày'), ('slot', 'Buổi'), ('student_code_snapshot', 'Mã học sinh'),
    ('full_name_snapshot', 'Họ tên'), ('class_code_snapshot', 'Mã lớp'), ('class_name_snapshot', 'Lớp'),
    ('grade_code_snapshot', 'Mã khối'), ('grade_name_snapshot', 'Khối'), ('status', 'Trạng thái'),
    ('source', 'Nguồn'), ('scanned_at', 'Giờ quét'), ('arrived_at', 'Giờ đến'),
    ('verified_at', 'Giờ xác minh'), ('arrival_status', 'Đúng giờ/muộn'),
    ('required_uniform_snapshot', 'Chuẩn đồng phục'), ('observed_uniform', 'Loại thực tế'),
    ('uniform_status', 'Kết quả đồng phục'), ('absence_reason', 'Lý do vắng'),
    ('uniform_reason', 'Lý do đồng phục'), ('uniform_note', 'Ghi chú'),
    ('checked_by_username', 'Người kiểm tra'), ('checked_at', 'Giờ kiểm tra'))
DATETIMES = {'scanned_at', 'arrived_at', 'verified_at', 'checked_at'}


def safe_cell(value):
    if not isinstance(value, str):
        return value
    # Spreadsheet parsers may ignore whitespace/control bytes before formulas.
    # XML 1.0 cannot represent these Unicode scalar values; remove them before
    # checking the resulting prefix so sanitization cannot expose a formula.
    cleaned = ''.join(c for c in value if c in '\t\n\r' or
                      (0x20 <= ord(c) <= 0xD7FF or 0xE000 <= ord(c) <= 0xFFFD or
                       0x10000 <= ord(c) <= 0x10FFFF))
    index = 0
    while index < len(cleaned) and (cleaned[index].isspace() or unicodedata.category(cleaned[index]) in ('Cc', 'Cf')):
        index += 1
    probe = cleaned[index:]
    risky = bool(probe) and probe[0] in '=+-@'
    return "'" + cleaned if risky else cleaned


def detail_values(row):
    result = []
    for key, _label in COLUMNS:
        value = row.get(key)
        if key in DATETIMES and value:
            value = datetime.fromisoformat(value).astimezone(ZoneInfo('Asia/Ho_Chi_Minh')).isoformat()
        result.append(safe_cell(value))
    return result


def export_response(result, format):
    headers = [label for _, label in COLUMNS]
    if format == 'csv':
        output = StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(detail_values(row) for row in result['items'])
        response = HttpResponse(output.getvalue().encode('utf-8-sig'), content_type='text/csv; charset=utf-8')
        # CSV remains one rectangular detail table; metadata accompanies download
        # in headers, and is also available from the detail JSON endpoint.
        response['X-Report-Timezone'] = 'Asia/Ho_Chi_Minh'
        response['X-Report-Generated-At'] = result['metadata']['generated_at']
        response['X-Report-Filters'] = json.dumps(result['metadata']['filters'], ensure_ascii=True)
        response['X-Report-Scope'] = json.dumps(result['metadata']['scope'], ensure_ascii=True)
    else:
        book = Workbook()
        details = book.active
        details.title = 'Chi_tiet'
        details.append(headers)
        for row in result['items']:
            details.append(detail_values(row))
        details.freeze_panes = 'A2'
        details.auto_filter.ref = details.dimensions
        summary = book.create_sheet('Tong_hop')
        for key, value in result['summary'].items():
            summary.append([key, value])
        metadata = book.create_sheet('Thong_tin')
        for key, value in result['metadata'].items():
            metadata.append([key, safe_cell(json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value)])
        output = BytesIO()
        book.save(output)
        response = HttpResponse(output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="attendance.{format}"'
    return response
